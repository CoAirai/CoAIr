"""Event Chronology Word export, following the COAIr Chronos template v1 (15 Oct 2026).

Rendered at download time from the stored result so that reviewer verification
and the latest layout apply to every chronology, including older ones.
"""

from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from .word_footnotes import xml_text

NAVY = "0C1E2E"
AMBER = "F8A810"
INK = "1A1A1A"
MUTE = "5B7386"
RULE = "D8DEE4"
ALT = "F5F8FA"
NOTE_BG = "FFF6E2"
DELAY_BG = "FFF1D6"
DELAY_INK = "9A5B00"
OK = "2E7D32"
WARN = "C62828"
SERIF = "Cambria"
SANS = "Calibri"
LOGO = Path(__file__).parent / "assets" / "coair_logo.png"

MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
MONTHS_LONG = ["January", "February", "March", "April", "May", "June", "July", "August",
               "September", "October", "November", "December"]

# Mirrors lib/chronology/events.ts so the Word file and the screen agree.
DOC_RULES: List[Tuple[str, str]] = [
    (r"\b(ins|ai|ei)-?\d|\binstruction\b", "Instruction"),
    (r"\bvo-?\d|\bvariation\b|\bchange (order|notice)\b|\bintc\b", "Variation"),
    (r"\bnot-?\d|\bnotice\b|\bnotification\b", "Notice"),
    (r"\bdet-?\d|\bdetermination\b", "Determination"),
    (r"\brfi-?\d|request for information", "RFI"),
    (r"\bncr-?\d|non-?conformance", "NCR"),
    (r"\bmom-?\d|\bminutes\b", "Minutes"),
    (r"\bmpr-?\d|progress report|period report", "Progress report"),
    (r"site diary|daily report", "Site diary"),
    (r"\bprg-?\d|\bprogramme\b|\bschedule\b", "Programme"),
    (r"\bspecification\b", "Specification"),
    (r"\bsubmittal\b", "Submittal"),
    (r"\bpurchase order\b", "Purchase order"),
    (r"\bregister\b|\btracker\b", "Register"),
    (r"\bdrawing\b|\.dwg\b", "Drawing"),
    (r"\.eml\b|\.msg\b|\bemail\b|\be-mail\b|^from:|\bsent:", "Email"),
    (r"\bletter\b|\bltr-?\d|\bdear\b", "Letter"),
    (r"payment cert|\bvaluation\b", "Payment certificate"),
    (r"\.jpe?g\b|\.png\b|\bphotograph", "Photograph"),
    (r"\bcontract\b|\bagreement\b|\bdeed\b", "Contract"),
    (r"\breport\b|\breview\b|\bpaper\b", "Report"),
]

DELAY_RULES: List[Tuple[str, str, List[str]]] = [
    ("N", "Notice", [r"\bnotice of delay\b", r"\bsub-?clause 20\.1\b",
                     r"\bextension of time\b.*\b(claim|submission|application)\b", r"(?-i:\bEoT\b)"]),
    ("I", "Instruction", [r"\binstruct(ed|ion)\b", r"\bvariation (order|instruction)\b",
                          r"\bchange of (scope|specification|finish)\b", r"\bsupersedes rev\b"]),
    ("H", "Hold", [r"\bsuspend(ed|sion)\b", r"\bstood down\b", r"\bon hold\b", r"\bawaiting approval\b",
                   r"\bwork(s)? stopped\b", r"\bheld pending\b"]),
    ("P", "Programme", [r"\bre-?sequenc", r"\brevised programme\b", r"\blead time\b",
                        r"\bagainst (the )?(baseline|plan(ned)?)\b", r"\bslipp(age|ed)\b"]),
    ("D", "Determination", [r"\bdetermin(es|ation|ed)\b", r"\bsub-?clause 3\.5\b",
                            r"\b(no|an) extension of time is (due|granted)\b"]),
    ("T", "Time language", [r"\bcritical (path|to)\b", r"\bprolongation\b",
                            r"\bdelay(ed|s)? (to|the) (completion|works|programme)\b", r"\btime[- ]related\b"]),
]

PARTIAL_REASONS = {
    "coverage_lost_to_budget": "Not every matching document could be read within the evidence budget for this run. Events held only in the unread documents are not recorded here.",
    "extraction_slice_skipped": "Parts of some long documents were not read. Events recorded only in those passages are not included.",
    "map_extraction_partial": "Event extraction did not complete for every document read, so some documents contributed fewer events than they may contain.",
    "counter_source_missing": "No records contradicting the entries were found. That may mean none exist, or that they are not in the knowledge base.",
}


@dataclass
class Source:
    source_id: str
    doc_id: str
    name: str
    title: str
    page: int
    pinpoint: str
    document_date: str


@dataclass
class Event:
    ref: str
    event_id: str
    date: str
    date_label: str
    inferred: bool
    text: str
    doc_type: str
    from_party: str
    to_party: str
    sources: List[Source]
    delay: Optional[Tuple[str, str]]
    reviewed: bool
    cited_ids: List[str] = field(default_factory=list)


def _doc_type(hint: str, kind: str, event_type: str) -> str:
    if kind == "email":
        return "Email"
    if kind == "excel":
        return "Register"
    lower = hint.lower()
    for pattern, label in DOC_RULES:
        if re.search(pattern, lower):
            return label
    if event_type and event_type.lower() not in ("event", "overview"):
        clean = re.sub(r"[_-]+", " ", event_type).strip()
        return clean[:1].upper() + clean[1:]
    return "Document"


def _delay(text: str) -> Optional[Tuple[str, str]]:
    for code, label, patterns in DELAY_RULES:
        if any(re.search(pattern, text, re.I) for pattern in patterns):
            return code, label
    return None


def format_date(raw: str, precision: str = "", long: bool = False) -> str:
    match = re.match(r"^(\d{4})(?:-(\d{2}))?(?:-(\d{2}))?", (raw or "").strip())
    if not match:
        return (raw or "").strip() or "Date not established"
    year, month, day = match.groups()
    names = MONTHS_LONG if long else MONTHS
    if precision == "year" or not month:
        return year
    month_name = names[int(month) - 1]
    if precision == "quarter":
        return f"Q{(int(month) + 2) // 3} {year}"
    if precision == "month" or not day:
        return f"{month_name} {year}"
    return f"{day} {month_name} {year}"


def _source(item: Dict[str, Any], index: int) -> Source:
    page = int(item.get("page") or 1)
    sheet, row_from, row_to = item.get("sheet"), item.get("row_from"), item.get("row_to")
    if sheet and row_from:
        pin = f"{sheet} row {row_from}" + (f"–{row_to}" if row_to and row_to != row_from else "")
    else:
        pin = f"p.{page}"
    return Source(
        source_id=str(item.get("source_id") or f"src_{index}"),
        doc_id=str(item.get("doc_id") or item.get("file_name") or ""),
        name=str(item.get("file_name") or item.get("title") or item.get("doc_id") or "Source"),
        title=str(item.get("title") or ""),
        page=page,
        pinpoint=pin,
        document_date=str(item.get("document_date") or ""),
    )


def build_events(result: Dict[str, Any]) -> Tuple[List[str], List[Event]]:
    evidence = result.get("evidence") or []
    raw_by_id = {str(e.get("source_id") or f"src_{i}"): e for i, e in enumerate(evidence)}
    source_by_id = {sid: _source(e, i) for i, (sid, e) in enumerate(raw_by_id.items())}
    reviews = result.get("reviews") or {}
    entries = result.get("entries") or []

    def texts(entry: Dict[str, Any]) -> List[Dict[str, Any]]:
        return [c for c in entry.get("claims") or [] if str(c.get("text") or "").strip()
                and c.get("supported", True)]

    overview = [str(c["text"]).strip() for e in entries
                if str(e.get("event_type") or "").lower() == "overview" for c in texts(e)]
    dated = [(i, e, str(e.get("event_date") or e.get("date") or "").strip())
             for i, e in enumerate(entries) if str(e.get("event_type") or "").lower() != "overview"]
    dated.sort(key=lambda item: (not item[2], item[2], item[0]))

    events: List[Event] = []
    for position, (index, entry, date) in enumerate(dated, 1):
        claims = texts(entry)
        text = " ".join(str(c["text"]).strip() for c in claims) or "Record not established from the documents reviewed."
        ids = list(dict.fromkeys(sid for c in claims for sid in c.get("source_ids") or [] if sid in source_by_id))
        sources = [source_by_id[sid] for sid in ids]
        primary = raw_by_id.get(ids[0]) if ids else {}
        primary = primary or {}
        hint = " ".join(filter(None, [str(primary.get("title") or ""), str(primary.get("file_name") or ""),
                                      str(primary.get("excerpt") or "")[:160]]))
        parties = [p for p in entry.get("parties") or [] if p]
        sender, recipient = primary.get("sender") or "", primary.get("recipient") or ""
        if sender or recipient:
            from_party, to_party = sender or "—", recipient or "—"
        elif len(parties) >= 2:
            from_party, to_party = parties[0], ", ".join(parties[1:])
        else:
            from_party, to_party = (parties[0] if parties else "—"), "—"
        ref = str(entry.get("entry_ref") or f"entry-{index}")
        inferred = "infer" in str(entry.get("date_source") or "").lower() or (
            bool(claims) and all(c.get("is_inference") for c in claims))
        events.append(Event(
            ref=ref,
            event_id=f"EV-{position:04d}",
            date=date,
            date_label=format_date(date, str(entry.get("date_precision") or "")) if date else "Date not established",
            inferred=inferred,
            text=text,
            doc_type=_doc_type(hint, str(primary.get("kind") or ""), str(entry.get("event_type") or "")),
            from_party=from_party,
            to_party=to_party,
            sources=sources,
            delay=_delay(text),
            reviewed=ref in reviews,
            cited_ids=ids,
        ))
    return overview, events


# ── low-level Word helpers ──────────────────────────────────────────────────

# Word rejects a file as unreadable when property children are out of schema order.
_TCPR_AFTER = {
    "w:tcBorders": ("w:shd", "w:noWrap", "w:tcMar", "w:textDirection", "w:tcFitText", "w:vAlign", "w:hideMark"),
    "w:shd": ("w:noWrap", "w:tcMar", "w:textDirection", "w:tcFitText", "w:vAlign", "w:hideMark"),
    "w:tcMar": ("w:textDirection", "w:tcFitText", "w:vAlign", "w:hideMark"),
}
_PPR_AFTER_BORDER = (
    "w:shd", "w:tabs", "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap", "w:overflowPunct",
    "w:topLinePunct", "w:autoSpaceDE", "w:autoSpaceDN", "w:bidi", "w:adjustRightInd", "w:snapToGrid",
    "w:spacing", "w:ind", "w:contextualSpacing", "w:mirrorIndents", "w:suppressOverlap", "w:jc",
    "w:textDirection", "w:textAlignment", "w:textboxTightWrap", "w:outlineLvl", "w:divId", "w:cnfStyle",
    "w:rPr", "w:sectPr", "w:pPrChange",
)


def _tcpr_insert(cell, element) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tag = element.tag.split("}")[1]
    existing = tc_pr.find(qn(f"w:{tag}"))
    if existing is not None:
        tc_pr.remove(existing)
    tc_pr.insert_element_before(element, *_TCPR_AFTER[f"w:{tag}"])


def _shade(cell, fill: str) -> None:
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), fill)
    _tcpr_insert(cell, shading)


def _cell_borders(cell, color: str = RULE, size: int = 2, left: Optional[Tuple[str, int]] = None,
                  none: bool = False) -> None:
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        element = OxmlElement(f"w:{edge}")
        if none:
            element.set(qn("w:val"), "nil")
        else:
            edge_color, edge_size = (left if edge == "left" and left else (color, size))
            element.set(qn("w:val"), "single")
            element.set(qn("w:sz"), str(edge_size))
            element.set(qn("w:color"), edge_color)
        borders.append(element)
    _tcpr_insert(cell, borders)


def _cell_margins(cell, top: int = 60, bottom: int = 60, side: int = 90) -> None:
    margins = OxmlElement("w:tcMar")
    for edge, value in (("top", top), ("bottom", bottom), ("left", side), ("right", side)):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")
        margins.append(element)
    _tcpr_insert(cell, margins)


def _repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    element = OxmlElement("w:tblHeader")
    element.set(qn("w:val"), "true")
    tr_pr.append(element)


def _cant_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement("w:cantSplit"))


def _para_border(paragraph, edge: str, color: str, size: int, space: int = 4) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    borders = p_pr.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        p_pr.insert_element_before(borders, *_PPR_AFTER_BORDER)
    element = OxmlElement(f"w:{edge}")
    element.set(qn("w:val"), "single")
    element.set(qn("w:sz"), str(size))
    element.set(qn("w:space"), str(space))
    element.set(qn("w:color"), color)
    borders.append(element)


def _field(paragraph, instruction: str, size: float, color: str) -> None:
    run = paragraph.add_run()
    _style(run, size=size, color=color)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    text = OxmlElement("w:instrText")
    text.set(qn("xml:space"), "preserve")
    text.text = instruction
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(text)
    run._r.append(end)


def _style(run, *, size: float = 10, color: str = INK, bold: bool = False, italic: bool = False,
           font: str = SANS, caps: bool = False) -> None:
    run.font.name = font
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), font)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold
    run.italic = italic
    run.font.all_caps = caps


def _run(paragraph, text: str, **style) -> None:
    run = paragraph.add_run(xml_text(text))
    _style(run, **style)


def _para(container, text: str = "", *, after: float = 7, before: float = 0, align=None, line: float = 1.15,
          **style):
    paragraph = container.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.line_spacing = line
    if align is not None:
        paragraph.alignment = align
    if text:
        _run(paragraph, text, **style)
    return paragraph


def _clear(cell):
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.0
    return paragraph


class _Builder:
    def __init__(self) -> None:
        self.doc = Document()
        normal = self.doc.styles["Normal"]
        normal.font.name = SANS
        normal.font.size = Pt(10)
        normal.font.color.rgb = RGBColor.from_string(INK)

    # text blocks
    def h1(self, text: str) -> None:
        p = _para(self.doc, text, before=13, after=7.5, size=15, bold=True, color=NAVY, font=SERIF)
        p.paragraph_format.keep_with_next = True
        _para_border(p, "bottom", AMBER, 10, 6)

    def h2(self, text: str) -> None:
        p = _para(self.doc, text, before=10, after=5, size=12, bold=True, color=NAVY, font=SERIF)
        p.paragraph_format.keep_with_next = True

    def body(self, text: str) -> None:
        _para(self.doc, text, after=7, line=1.17)

    def bullet(self, text: str) -> None:
        p = _para(self.doc, after=4.5, line=1.17)
        p.paragraph_format.left_indent = Cm(0.63)
        p.paragraph_format.first_line_indent = Cm(-0.4)
        _run(p, "•  ", color=AMBER, bold=True)
        _run(p, text)

    def spacer(self, after: float = 8) -> None:
        _para(self.doc, after=after)

    # tables
    def table(self, widths_cm: List[float], header: Optional[List[str]], rows: List[List[Any]],
              *, zebra: bool = True, borders: bool = True, center: Tuple[int, ...] = ()):
        table = self.doc.add_table(rows=0, cols=len(widths_cm))
        table.alignment = WD_TABLE_ALIGNMENT.LEFT
        table.autofit = False
        if header:
            row = table.add_row()
            _repeat_header(row)
            for index, label in enumerate(header):
                cell = row.cells[index]
                cell.width = Cm(widths_cm[index])
                _shade(cell, NAVY)
                _cell_borders(cell)
                _cell_margins(cell, 80, 80)
                p = _clear(cell)
                if index in center:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                _run(p, label, size=7.5, bold=True, color="FFFFFF", caps=True)
        for number, values in enumerate(rows):
            row = table.add_row()
            _cant_split(row)
            for index, value in enumerate(values):
                cell = row.cells[index]
                cell.width = Cm(widths_cm[index])
                spec = value if isinstance(value, dict) else {"text": value}
                fill = spec.get("fill") or (ALT if zebra and number % 2 else None)
                if fill:
                    _shade(cell, fill)
                if borders:
                    _cell_borders(cell, left=spec.get("left"))
                else:
                    _cell_borders(cell, none=True)
                _cell_margins(cell, spec.get("pad", 60), spec.get("pad", 60))
                lines = spec.get("lines") or [spec]
                for line_no, line in enumerate(lines):
                    p = _clear(cell) if line_no == 0 else cell.add_paragraph()
                    p.paragraph_format.space_after = Pt(0)
                    p.paragraph_format.line_spacing = 1.0
                    if index in center or spec.get("center"):
                        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for part in line.get("runs") or [line]:
                        _run(p, str(part.get("text") or ""), size=part.get("size", 8.5),
                             bold=part.get("bold", False), color=part.get("color", INK),
                             italic=part.get("italic", False), font=part.get("font", SANS))
        return table

    # page furniture
    def portrait(self, section) -> None:
        section.orientation = WD_ORIENT.PORTRAIT
        section.page_width, section.page_height = Cm(21.0), Cm(29.7)
        section.left_margin = section.right_margin = Cm(2.0)
        section.top_margin, section.bottom_margin = Cm(2.47), Cm(2.12)
        section.header_distance, section.footer_distance = Cm(0.99), Cm(0.74)

    def landscape(self, section) -> None:
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width, section.page_height = Cm(29.7), Cm(21.0)
        section.left_margin = section.right_margin = Cm(1.5)
        section.top_margin, section.bottom_margin = Cm(2.47), Cm(1.94)
        section.header_distance, section.footer_distance = Cm(0.99), Cm(0.71)

    def furniture(self, section, *, title: str, wide: bool) -> None:
        section.header.is_linked_to_previous = False
        section.footer.is_linked_to_previous = False
        header = section.header.paragraphs[0]
        header.style = self.doc.styles["Normal"]
        header.paragraph_format.space_after = Pt(2)
        width = section.page_width - section.left_margin - section.right_margin
        header.paragraph_format.tab_stops.add_tab_stop(width, WD_TAB_ALIGNMENT.RIGHT)
        if LOGO.is_file():
            header.add_run().add_picture(str(LOGO), width=Cm(1.3))
        _run(header, "\t")
        _run(header, "Event Chronology", size=8, bold=True, color=NAVY)
        _run(header, "   |   ", size=8, color=RULE)
        _run(header, title, size=8, color=MUTE)
        _para_border(header, "bottom", AMBER, 8, 2)

        footer = section.footer.paragraphs[0]
        footer.style = self.doc.styles["Normal"]
        footer.paragraph_format.space_after = Pt(0)
        footer.paragraph_format.tab_stops.add_tab_stop(width, WD_TAB_ALIGNMENT.RIGHT)
        _para_border(footer, "top", RULE, 4, 4)
        _run(footer, "Generated by COAIr · Chronology module  ", size=7, color=MUTE)
        _run(footer, "Preliminary output — subject to review by the instructed professional", size=7,
             color=MUTE, italic=True)
        _run(footer, "\tPage ", size=7, color=MUTE)
        _field(footer, "PAGE", 7, MUTE)
        _run(footer, " of ", size=7, color=MUTE)
        _field(footer, "NUMPAGES", 7, MUTE)


def _plural(count: int, word: str, plural: str = "") -> str:
    return f"{count} {word if count == 1 else (plural or word + 's')}"


def build_event_chronology_docx(*, job: Dict[str, Any], project_name: str,
                                generated: Optional[datetime] = None) -> bytes:
    result = job.get("result") or {}
    request = job.get("request") or {}
    topic = str(request.get("topic") or job.get("title") or "Chronology")
    overview, events = build_events(result)
    generated = generated or datetime.now(timezone.utc)
    issue_date = f"{generated.day:02d} {MONTHS_LONG[generated.month - 1]} {generated.year}"
    sequence = job.get("sequence_number") or 1
    reference = f"CHR-{int(sequence):03d}"
    verified = sum(1 for e in events if e.reviewed)
    unverified = len(events) - verified
    delay_events = [e for e in events if e.delay]
    dated = [e for e in events if e.date]
    period = (f"{format_date(dated[0].date, long=True)} – {format_date(dated[-1].date, long=True)}"
              if dated else "Not established from the documents reviewed")
    cited_docs: Dict[str, Tuple[Source, Event]] = {}
    for event in events:
        for source in event.sources:
            cited_docs.setdefault(source.doc_id, (source, event))
    considered = result.get("research_audit", {}).get("candidates") if isinstance(
        result.get("research_audit"), dict) else None
    considered_count = len(considered) if isinstance(considered, list) else None
    status = "ISSUED — all entries verified" if events and not unverified else "PRELIMINARY — for internal review"
    short_title = f"{topic[:70]}  ·  {project_name[:50]}"

    b = _Builder()
    doc = b.doc

    # ── cover ───────────────────────────────────────────────────────────────
    cover = doc.sections[0]
    b.portrait(cover)
    cover.footer.is_linked_to_previous = False
    footer = cover.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _run(footer, "CO Consultants FZ  ·  thecoconsultants.com  ·  coair.ai", size=7, color=MUTE)
    if LOGO.is_file():
        logo = _para(doc, before=30, after=30)
        logo.add_run().add_picture(str(LOGO), width=Cm(3.0))
    _para(doc, "Chronology module  ·  Preliminary output", after=4.5, size=7.5, bold=True, color=AMBER, caps=True)
    _para(doc, "Event Chronology", after=3, size=26, bold=True, color=NAVY, font=SERIF)
    subtitle = _para(doc, topic, after=14, size=13, color=MUTE, font=SERIF)
    _para_border(subtitle, "bottom", AMBER, 16, 8)
    meta = [
        ("Project", project_name),
        ("Contract", "[Contract reference and form]"),
        ("Prepared for", "[Instructing party]"),
        ("Prepared by", "[Analyst name], CO Consultants FZ"),
        ("Knowledge base", f"{project_name}" + (f" ({considered_count:,} documents considered for this topic)"
                                                 if considered_count else "")),
        ("Period covered", period),
        ("Events recorded", f"{len(events)}" + (f" ({events[0].event_id} to {events[-1].event_id})" if events else "")),
        ("Document reference", reference),
        ("Revision", "A — first issue"),
        ("Date of issue", issue_date),
        ("Status", status),
    ]
    b.table([5.2, 11.8], None, [
        [{"text": k, "bold": True, "color": NAVY, "pad": 70}, {"text": v, "pad": 70}] for k, v in meta
    ], borders=False)
    b.spacer(14)
    b.table([17.0], None, [[{
        "fill": NOTE_BG, "pad": 140, "left": (AMBER, 18),
        "lines": [
            {"text": "CONFIDENTIAL", "bold": True, "color": NAVY, "size": 8},
            {"text": "Prepared in contemplation of a contractual claim. This document is a preliminary, "
                     "machine-assisted compilation of the project record. It records what the documents say; "
                     "it does not state an opinion on entitlement, causation or liability.", "size": 8},
        ],
    }]], zebra=False)

    # ── front matter (portrait) ─────────────────────────────────────────────
    front = doc.add_section(WD_SECTION.NEW_PAGE)
    b.portrait(front)
    b.furniture(front, title=short_title, wide=False)
    b.h1("1.  Basis and limitations")
    b.body("This chronology was compiled by the COAIr Chronology module from documents held in the project "
           "knowledge base identified on the cover page. Each entry is traceable to source documents and a "
           "pinpoint location within them.")
    b.h2("1.1  What this document is")
    b.bullet("A factual record of events, in date order, as evidenced by contemporaneous project documents.")
    b.bullet("A citation index — every entry identifies the source document, page and passage from which it was derived.")
    b.bullet("A first pass intended to be checked, corrected and extended by the instructed professional.")
    b.h2("1.2  What this document is not")
    b.bullet("It is not a delay analysis. No critical path has been identified and no method under the SCL "
             "Protocol or AACE RP 29R-03 has been applied.")
    b.bullet("It does not attribute responsibility for any event, or assess whether any event caused delay to completion.")
    b.bullet("It is not expert evidence and must not be issued as such without review, verification and "
             "adoption by a named professional.")
    b.h2("1.3  Machine-assisted preparation")
    b.body("Entries were extracted automatically and are marked unverified until a reviewer confirms them "
           "against the source. Automated extraction can misread dates, conflate similar documents and miss "
           "events held only in formats the knowledge base cannot index. The reviewer's verification, not the "
           "extraction, is what makes an entry reliable; verification status for this issue is at section 2.3. "
           "Where this chronology is relied upon in a submission, state that AI-assisted tooling was used and "
           "describe the validation performed.")
    if overview:
        b.h2("1.4  Overview of the record")
        for text in overview:
            b.bullet(text)

    b.h1("2.  Records reviewed")
    b.h2("2.1  Scope")
    scope = (f"The topic searched was “{topic}”. "
             + (f"{considered_count:,} documents in the knowledge base were considered for this topic, and "
                if considered_count else "")
             + f"{_plural(len(cited_docs), 'document')} {'is' if len(cited_docs) == 1 else 'are'} cited in "
               f"section 4. The cited documents are summarised by record type below.")
    b.body(scope)
    by_type: Dict[str, int] = {}
    for source, event in cited_docs.values():
        by_type[event.doc_type] = by_type.get(event.doc_type, 0) + 1
    b.table([8.0, 2.5, 6.5], ["Record type", "Count", "Completeness"], [
        [kind, {"text": str(count), "center": True}, {"text": "Not assessed", "color": MUTE}]
        for kind, count in sorted(by_type.items(), key=lambda item: (-item[1], item[0]))
    ] or [["No documents cited", {"text": "0", "center": True}, "—"]], center=(1,))
    b.h2("2.2  Records not available")
    b.body("The following limitations were identified. They are stated here because an absent record is itself "
           "a finding — it limits what any subsequent analysis can conclude, and it may be recoverable on request.")
    reasons = [PARTIAL_REASONS[r] for r in result.get("partial_reasons") or [] if r in PARTIAL_REASONS]
    undated = [e for e in events if not e.date]
    if undated:
        reasons.append(f"{_plural(len(undated), 'entry', 'entries')} could not be dated from the documents "
                       f"({', '.join(e.event_id for e in undated)}) and {'is' if len(undated) == 1 else 'are'} "
                       "placed at the end of section 4.")
    for reason in reasons or ["The module identified no gaps in the records it read. The reviewer should "
                              "confirm whether any expected record types are absent from the knowledge base."]:
        b.bullet(reason)
    b.h2("2.3  Verification status of this issue")
    b.body(f"Of {_plural(len(events), 'entry', 'entries')}, {verified} {'has' if verified == 1 else 'have'} been "
           f"verified against source by a reviewer and {unverified} {'remains' if unverified == 1 else 'remain'} "
           "unverified. Unverified entries are marked U in the Rec. column and must be checked before this "
           "chronology is relied upon externally.")

    b.h1("3.  Conventions")
    b.h2("3.1  Event identifiers")
    b.body("Events are numbered EV-0001 upward in date order at first issue. An identifier, once allocated, is "
           "never reused or reassigned. Entries that cannot be dated from the documents are numbered last.")
    b.h2("3.2  Dates")
    b.body("Dates are given to the precision the source supports: a day, a month or a year. A date shown in "
           "square brackets is inferred from context rather than stated in the source.")
    b.h2("3.3  Record status codes")
    b.table([1.6, 3.6, 11.8], ["Code", "Meaning", "Applied when"], [
        [{"text": "C", "bold": True, "color": OK, "center": True}, "Contemporaneous",
         "Recorded at or about the time of the event, by a party with direct knowledge."],
        [{"text": "R", "bold": True, "color": DELAY_INK, "center": True}, "Retrospective",
         "Recorded materially after the event. The interval is stated in the entry."],
        [{"text": "I", "bold": True, "color": DELAY_INK, "center": True}, "Inferred",
         "The event itself is not directly recorded; it is inferred from other documents."],
        [{"text": "U", "bold": True, "color": WARN, "center": True}, "Unverified",
         "Extracted automatically and not yet checked against source by a reviewer."],
    ], center=(0,))
    b.spacer(4)
    b.h2("3.4  Delay markers")
    b.body("Rows shaded amber with a ▲ contain delay-related wording — a notice, instruction, hold, programme "
           "change, determination or time language. The marker describes what the record says. It is not a "
           "finding that delay occurred, that it affected completion, or who is responsible.")
    b.h2("3.5  Source references")
    b.body("Each entry cites the source document reference followed by a pinpoint — page and, where the document "
           "is structured, sheet and row. The same document can be opened at that page in COAIr. Readers without "
           "COAIr access should use Appendix A, which gives the full reference for every document cited so that "
           "it can be located in the disclosure bundle.")

    # ── chronology (landscape) ──────────────────────────────────────────────
    wide = doc.add_section(WD_SECTION.NEW_PAGE)
    b.landscape(wide)
    b.furniture(wide, title=short_title, wide=True)
    b.h1("4.  Chronology of events")
    b.body(f"{_plural(len(events), 'event')}, in date order"
           + (f", for the period {period}" if dated else "")
           + ". Each entry states what the record says and where that record is. Entries carrying a U in the "
             "Rec. column have not yet been verified against source.")
    rows = []
    for number, event in enumerate(events):
        fill = DELAY_BG if event.delay else (ALT if number % 2 else None)
        refs = [{"runs": [
            {"text": s.name.rsplit(".", 1)[0], "bold": True, "color": "1F5C8B", "size": 8.5},
            {"text": f"  {s.pinpoint}", "color": MUTE, "size": 8},
        ]} for s in event.sources[:4]] or [{"text": "—", "color": MUTE}]
        if len(event.sources) > 4:
            refs.append({"text": f"+{len(event.sources) - 4} more (Appendix A)", "color": MUTE, "size": 7.5})
        status = ["I" if event.inferred else "C"] + ([] if event.reviewed else ["U"])
        event_lines = [{"text": event.text}]
        if event.delay:
            event_lines.append({"text": f"▲ {event.delay[1]}", "bold": True, "color": DELAY_INK, "size": 8})
        rows.append([
            {"text": event.event_id, "bold": True, "color": NAVY, "fill": fill,
             "left": (AMBER, 24) if event.delay else None},
            {"text": f"[{event.date_label}]" if event.inferred else event.date_label, "bold": True,
             "fill": fill, "color": DELAY_INK if event.inferred else INK},
            {"lines": event_lines, "fill": fill},
            {"text": event.doc_type, "color": MUTE, "fill": fill},
            {"lines": refs, "fill": fill},
            {"lines": [{"text": event.from_party, "size": 8}, {"text": f"→ {event.to_party}", "size": 8,
                                                                 "color": MUTE}], "fill": fill},
            {"lines": [{"text": code, "bold": True, "color": {"C": OK, "I": DELAY_INK, "U": WARN}[code]}
                       for code in status], "fill": fill, "center": True},
        ])
    b.table([1.9, 2.3, 10.4, 2.3, 4.5, 3.9, 1.4],
            ["Event ID", "Date", "Event", "Record type", "Source reference", "From / To", "Rec."],
            rows or [["—", "—", "No dated events were returned.", "", "", "", ""]], center=(6,))

    b.h1("5.  Delay event schedule")
    note = _para(doc, after=8, line=1.1)
    note.paragraph_format.left_indent = Cm(0.3)
    _para_border(note, "left", AMBER, 18, 8)
    _run(note, "Analyst entry — not generated by COAIr.  ", size=8.5, bold=True, color=DELAY_INK)
    _run(note, "The columns below call for judgement on criticality, impact and concurrency. COAIr does not "
               "populate them. They are provided so that the analyst's assessment sits alongside, and remains "
               "visibly distinct from, the factual record above.", size=8.5)
    b.body("Add a row for each event the analyst considers may bear on completion. Reference the Event ID; do not "
           "restate the facts, which are above. Entries marked ▲ are pre-listed.")
    listed = delay_events or events[:3]
    b.table([2.4, 3.6, 2.8, 2.8, 2.8, 12.3],
            ["Event ID", "Criticality", "Delay (days)", "Notice given", "Concurrent", "Analyst comment"],
            [[{"text": e.event_id, "bold": True, "color": NAVY}, "", "", "", "", ""] for e in listed]
            + [["", "", "", "", "", ""]] * max(0, 3 - len(listed)), center=(2, 3, 4))

    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    b.h1("Appendix A  —  Source register")
    b.body("Every document cited in section 4, with its full reference. Where a reader has no access to COAIr, "
           "this register is what allows each entry to be located in the disclosure bundle.")
    register = []
    for source, event in sorted(cited_docs.values(), key=lambda item: item[0].name.lower()):
        cited_at = [e.event_id for e in events if any(s.doc_id == source.doc_id for s in e.sources)]
        register.append([
            {"text": source.name.rsplit(".", 1)[0], "bold": True, "color": "1F5C8B"},
            format_date(source.document_date) if re.match(r"^\d{4}", source.document_date)
            else (source.document_date or "—"),
            source.title or "—",
            {"text": event.doc_type, "color": MUTE},
            {"text": source.doc_id[:12], "color": MUTE, "font": "Consolas", "size": 7.5},
            {"text": ", ".join(cited_at), "bold": True, "color": NAVY},
        ])
    b.table([4.4, 2.6, 9.4, 2.8, 3.2, 4.3],
            ["Document ref", "Date", "Title", "Type", "KB document ID", "Cited at"],
            register or [["—", "", "No documents cited", "", "", ""]])
    b.spacer(12)
    end = _para(doc, after=0)
    _para_border(end, "top", RULE, 6, 6)
    _run(end, "End of chronology.  ", size=8.5, bold=True, color=NAVY)
    _run(end, f"{reference} Rev A  ·  Generated {issue_date}  ·  {_plural(len(events), 'event')}  ·  "
              f"{_plural(len(cited_docs), 'source document')}  ·  {verified} of {len(events)} entries verified.",
         size=8.5, color=MUTE)

    doc.core_properties.title = f"Event Chronology — {topic}"
    doc.core_properties.author = "CO Consultants FZ — COAIr"
    raw = io.BytesIO()
    doc.save(raw)
    return raw.getvalue()


__all__ = ["build_event_chronology_docx", "build_events", "format_date"]
