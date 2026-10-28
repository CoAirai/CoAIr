"""Format caller-owned answer rows as Excel; never fetch project data or run SQL."""
from io import BytesIO
import math
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

MAX_ROWS = 2000
MAX_COLUMNS = 60
MAX_CELL_CHARS = 32767


def build_answer_xlsx(columns, rows):
    if not columns or len(columns) > MAX_COLUMNS or len(rows) > MAX_ROWS:
        raise ValueError("Export supports 1–60 columns and up to 2000 retained rows.")
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Answer results"
    for values in [columns, *rows]:
        if len(values) > len(columns):
            raise ValueError("A row has more values than the supplied columns.")
        sheet.append([None] * len(columns))
        for column, value in enumerate(values[:len(columns)], 1):
            cell = sheet.cell(sheet.max_row, column)
            if value is None or isinstance(value, (bool, int, float)):
                if isinstance(value, float) and not math.isfinite(value):
                    raise ValueError("Excel cannot represent non-finite numbers.")
                # Excel cannot faithfully represent integers with more than 15 digits.
                if isinstance(value, int) and not isinstance(value, bool) and len(str(abs(value))) > 15:
                    cell.value = str(value)
                    cell.data_type = "s"
                else:
                    cell.value = value
            else:
                text = str(value)
                if ILLEGAL_CHARACTERS_RE.search(text):
                    raise ValueError("A cell contains characters Excel cannot represent.")
                if len(text) > MAX_CELL_CHARS:
                    raise ValueError("A cell exceeds Excel's 32767 character limit.")
                cell.value = text
                # Treat formula-looking source text as literal, including hyperlinks.
                cell.data_type = "s"
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column in range(1, len(columns) + 1):
        sheet.column_dimensions[get_column_letter(column)].width = 24
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
