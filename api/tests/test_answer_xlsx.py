from io import BytesIO
import pytest
from openpyxl import load_workbook
from src.answer_xlsx import build_answer_xlsx


def test_workbook_preserves_values_and_disables_formulas():
    workbook = load_workbook(BytesIO(build_answer_xlsx(["Zero", "Approved", "Source", "ID"],
        [[0, False, '=HYPERLINK("https://example.com")', 1234567890123456789]])))
    sheet = workbook.active
    assert sheet["A2"].value == 0
    assert sheet["B2"].value is False
    assert sheet["C2"].value.startswith("=HYPERLINK")
    assert sheet["C2"].data_type == "s"
    assert sheet["D2"].value == "1234567890123456789"
    assert sheet.freeze_panes == "A2"
    assert sheet.auto_filter.ref == "A1:D2"
    assert sheet["A1"].font.bold


def test_empty_rows_keep_headers_and_limits_reject_instead_of_truncate():
    sheet = load_workbook(BytesIO(build_answer_xlsx(["Amount"], []))).active
    assert sheet["A1"].value == "Amount"
    with pytest.raises(ValueError):
        build_answer_xlsx(["A"], [[0]] * 2001)
    with pytest.raises(ValueError):
        build_answer_xlsx(["A"], [["x" * 32768]])


def test_compact_headings_remain_bold_in_word():
    from docx import Document
    from src.answer_docx import build_answer_docx
    from src.markdown_blocks import parse
    assert parse("###Summary") == [{"type":"heading", "level":3, "text":"Summary"}]
    assert parse("#hashtag")[0]["type"] == "para"
    document = Document(BytesIO(build_answer_docx("", "###Summary\n\n**Evidence**\n\n| A | B |\n| --- | --- |\n| 0 | false |")))
    heading = next(p for p in document.paragraphs if p.text == "Summary")
    assert heading.runs[0].font.bold
    assert document.tables[0].cell(1, 0).text == "0"


@pytest.mark.parametrize("rows", [[[float("nan")]], [["bad\x00value"]], [[1, 2]]])
def test_invalid_cells_and_ragged_overflow_are_rejected(rows):
    with pytest.raises(ValueError):
        build_answer_xlsx(["A"], rows)


def test_export_http_requires_authentication_and_formats_caller_owned_data():
    from fastapi import FastAPI
    from fastapi.testclient import TestClient
    from backend.api.chat import router
    from backend.core.security import get_current_user
    app = FastAPI()
    app.include_router(router)
    with TestClient(app) as client:
        assert client.post("/chat/spreadsheet", json={"table_columns":["A"], "table_rows":[[0]]}).status_code == 401
        app.dependency_overrides[get_current_user] = lambda: object()
        result = client.post("/chat/spreadsheet", json={"table_columns":["A"], "table_rows":[[0]]})
        assert result.status_code == 200
        assert result.headers["content-type"].endswith("spreadsheetml.sheet")
        assert load_workbook(BytesIO(result.content)).active["A2"].value == 0
        invalid = client.post("/chat/spreadsheet", json={"table_columns":[], "table_rows":[]})
        assert invalid.status_code == 422
