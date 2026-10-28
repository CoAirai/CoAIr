"""Long evidence remains eligible while individual model inputs stay bounded."""
from types import SimpleNamespace
from unittest.mock import patch
import io

import duckdb
from docx import Document

from backend.services.chat_orchestrator import ChatOrchestrator
from src import ai_reports, chronology_v2
from src.evidence_model import EvidenceItem
from src.evidence_pack import select_pack


def test_selected_chat_can_read_late_pages_and_late_selected_files():
    con = duckdb.connect(':memory:')
    con.execute('CREATE TABLE chunks(project_id VARCHAR, doc_id VARCHAR, file_name VARCHAR, page_number INTEGER, chunk_id VARCHAR, text VARCHAR)')
    records = [SimpleNamespace(doc_id=f'd{i}', project_id='p', file_name=f'file{i}.pdf', file_type='pdf') for i in range(30)]
    records += [SimpleNamespace(doc_id='foreign', project_id='other', file_name='file29.pdf', file_type='pdf')]
    for rec in records:
        con.execute('INSERT INTO chunks VALUES (?,?,?,?,?,?)', [rec.project_id,rec.doc_id,rec.file_name,1,'a','background text ' * 3000])
    con.execute('INSERT INTO chunks VALUES (?,?,?,?,?,?)', ['p','d29','file29.pdf',200,'late','Zirconium compensation was approved by the engineer.'])
    con.execute('INSERT INTO chunks VALUES (?,?,?,?,?,?)', ['other','foreign','file29.pdf',201,'secret','Zirconium SECRET FOREIGN PRICE'])
    registry = SimpleNamespace(get=lambda doc_id: next((r for r in records if r.doc_id == doc_id), None), get_all=lambda project_id: [r for r in records if r.project_id == project_id])
    try:
        with patch('src.document_registry.get_document_registry', return_value=registry), patch('src.chunk_store.get_chunk_store', return_value=SimpleNamespace(connection=lambda: con)):
            context = ChatOrchestrator.__new__(ChatOrchestrator)._build_document_context(
                [r.doc_id for r in records], 'p', query='What zirconium compensation was approved?')
        assert 'Zirconium compensation was approved' in context
        assert 'Page: 200' in context
        assert 'SECRET' not in context
        assert len(context) < 16500
        assert 'not a complete reading' in context
    finally:
        con.close()


def test_same_filename_is_not_an_ambiguous_legacy_identity():
    from src.selected_passages import selected_passages
    con = duckdb.connect(':memory:')
    con.execute('CREATE TABLE chunks(project_id VARCHAR, doc_id VARCHAR, file_name VARCHAR, page_number INTEGER, chunk_id VARCHAR, text VARCHAR)')
    con.execute("INSERT INTO chunks VALUES ('p','legacy','same.pdf',1,'c','ambiguous content')")
    registry = SimpleNamespace(get_all=lambda **kw: [SimpleNamespace(doc_id='a', file_name='same.pdf'), SimpleNamespace(doc_id='b',file_name='same.pdf')])
    try:
        assert list(selected_passages(con,registry,'p','a','same.pdf')) == []
    finally:
        con.close()


def test_full_report_pack_keeps_late_facts_even_with_repeated_dates():
    evidence = [EvidenceItem(f's{i}','d','large.pdf',page=i+1,excerpt='3 May 2006 notice ' + str(i) + 'x'*3400,score=1) for i in range(180)]
    selection = select_pack(evidence,max_chars=None,relevance_floor=0)
    assert selection.evidence == evidence
    assert selection.stats['selected_chars'] > 600000
    assert selection.stats['dropped_budget'] == selection.stats['dropped_document_cap'] == 0
    batches = chronology_v2.evidence_batches(selection.evidence)
    assert len(batches) > 3
    assert [e.source_id for batch in batches for e in batch] == [e.source_id for e in evidence]
    assert all(len(chronology_v2.evidence_markdown(batch)) <= chronology_v2.EVIDENCE_BATCH_CHARS for batch in batches)


def test_extraction_accepts_more_than_eighteen_material_events():
    event = {'claims':[{'text':'Notice issued','source_ids':['s1']}]}
    assert len(chronology_v2.ExtractionModel.model_validate({'entries':[event]*40}).entries) == 40


def test_selected_long_pack_reaches_word_and_retains_last_event(monkeypatch):
    evidence = [EvidenceItem(f's{i}',f'd{i%3}',f'file{i%3}.pdf',page=i+1,excerpt=f'Engineer issued notice {i}. '+'x'*3400) for i in range(100)]
    events = [{'claims':[{'text':f'Engineer issued notice {i}.','source_ids':[e.source_id]}]} for i,e in enumerate(evidence)]
    monkeypatch.setattr(chronology_v2,'evidence_from_documents',lambda *a: evidence)
    def extract(items, *a, **kw):
        assert items == evidence
        return events
    monkeypatch.setattr(chronology_v2,'extract_batches',extract)
    monkeypatch.setattr(chronology_v2,'synthesize',lambda **kw: {'overview_claims':events[0]['claims'],'entries':events})
    monkeypatch.setattr(chronology_v2,'verify_claims',lambda **kw: {'overview:0': True, **{f'event:{i}:0':True for i in range(100)}})
    monkeypatch.setattr(ai_reports,'_model_for_task',lambda *a:'qualified-model')
    result = ai_reports._generate_chronology_v2(project_id='p',project_name='Long evidence',topic='Notices',source_doc_ids=['d0','d1','d2'],preparation={'prepared':{'original_query':'Notices','english_query':'Notices'}})
    assert result['pack']['chars'] > 240000
    assert result['audit']['issued_events'] == 100
    document = Document(io.BytesIO(result['docx']))
    text = '\n'.join(p.text for p in document.paragraphs) + '\n'.join(c.text for t in document.tables for r in t.rows for c in r.cells)
    assert 'Engineer issued notice 99.' in text


def test_forensic_reads_all_batches_including_oversized_source(monkeypatch):
    import json
    from src import llm_client
    evidence = [EvidenceItem(f's{i}',f'd{i}',f'f{i}.pdf',excerpt=f'Engineer issued notice {i}. '+'x'*3400) for i in range(100)]
    evidence.append(EvidenceItem('large','large','large.pdf',excerpt='background '*16000+'Engineer approved zirconium compensation.'))
    seen = []
    def generate(prompt, **kwargs):
        payload = prompt.split('EVIDENCE JSON:\n',1)[1]
        assert len(payload) <= 80000
        batch = json.loads(payload)
        seen.extend(batch)
        claims = [{'text': 'Engineer approved zirconium compensation.', 'source_ids':['large'], 'counter_source_ids':[], 'confidence':'high','missing_records':[]}] if any('zirconium' in item['excerpt'] for item in batch) else []
        return SimpleNamespace(raw={'sections':[{'name':'Factual chronology','claims':claims}]},usage=SimpleNamespace(model='qualified-model'))
    monkeypatch.setattr(ai_reports,'build_research_plan',lambda *a:['q'])
    monkeypatch.setattr(ai_reports,'retrieve_evidence',lambda *a:evidence)
    monkeypatch.setattr(llm_client,'generate_response_json',generate)
    monkeypatch.setattr('backend.tasks.report_jobs.get_report_job_store',lambda:SimpleNamespace(list_project=lambda *a,**k:[]))
    result = ai_reports.generate_forensic(project_id='p',project_name='Qualification',topic='Compensation')
    assert result['evidence_batches'] > 3
    assert {item['source_id'] for item in seen} == {e.source_id for e in evidence}
    assert result['sections']['Factual chronology'][0]['text'] == 'Engineer approved zirconium compensation.'
    assert result['docx'].startswith(b'PK')


def test_overview_reduction_never_reinflates_the_full_report_context(monkeypatch):
    import json
    from src import llm_client
    evidence = [EvidenceItem(f's{i}',f'd{i}',f'f{i}.pdf',excerpt=f'Engineer issued notice {i}. '+'x'*5900) for i in range(100)]
    candidates = [{'claims':[{'text':f'Engineer issued notice {i}.','source_ids':[e.source_id]}]} for i,e in enumerate(evidence)]
    calls = []
    def generate(prompt, **kwargs):
        assert len(prompt) < 190000
        events = json.loads(prompt.split('EVENT CLUSTERS:\n',1)[1].split('\n\nEVIDENCE BEGIN',1)[0])
        calls.append(kwargs['schema_name'])
        if kwargs['schema_name'] == 'chronology_v2_overview':
            raw = {'overview_claims':[events[0]['claims'][0]]}
        else:
            raw = {'entries':events}
        return SimpleNamespace(raw=raw)
    monkeypatch.setattr(llm_client,'generate_response_json',generate)
    prepared = chronology_v2.PreparedChronologyQuery('notices','notices',(),(),(),(),(),('q',))
    draft = chronology_v2.synthesize(prepared=prepared,candidates=candidates,evidence=evidence,cache_context='long-qualification')
    assert len(draft['entries']) == 100
    assert calls.count('chronology_v2_overview') > 1
    assert {e['claims'][0]['source_ids'][0] for e in draft['entries']} == {e.source_id for e in evidence}


def test_report_selection_has_no_five_hundred_document_ceiling():
    from backend.api.reports import ChronologyGenerateRequest
    request = ChronologyGenerateRequest(topic='All selected notices',source_doc_ids=[f'd{i}' for i in range(1000)])
    assert len(request.source_doc_ids) == 1000
