import io
import zipfile
from datetime import datetime, timezone
from xml.etree import ElementTree

from docx import Document
from docx.enum.section import WD_ORIENT

from src.chronology_template_docx import build_event_chronology_docx, build_events, format_date

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"

JOB = {
    "job_id": "job-1",
    "title": "utility diversion",
    "sequence_number": 9,
    "request": {"topic": "utility diversion", "project_name": "Edinburgh Tram Inquiry (public)"},
    "result": {
        "entries": [
            {"entry_ref": "6.9.1", "event_type": "overview",
             "claims": [{"text": "MUDFA diversions ran late across the network.", "source_ids": ["src_a"]}]},
            {"entry_ref": "6.9.3", "event_date": "2007-11-25", "date_precision": "exact",
             "parties": ["tie limited", "Infraco"],
             "claims": [{"text": "Infraco gave notice of delay for the Gogar works.", "source_ids": ["src_a", "src_b"]}]},
            {"entry_ref": "6.9.2", "event_date": "2005-09-19", "date_precision": "exact",
             "claims": [{"text": "tie and Parsons Brinckerhoff executed the SDS Agreement.", "source_ids": ["src_b"]}]},
            {"entry_ref": "6.9.4", "event_date": "", "date_precision": "unknown",
             "claims": [{"text": "Undated record \x0b with a control character.", "source_ids": []}]},
        ],
        "evidence": [
            {"source_id": "src_a", "doc_id": "62356948e4b30cf7", "file_name": "CEC00793598.pdf",
             "title": "ETN INFRACO NOTIFICATION of tie CHANGE (INTC)", "document_date": "2008-05-09", "page": 2},
            {"source_id": "src_b", "doc_id": "8bacd3ccdeb70b9d", "file_name": "PBH00035962.pdf",
             "title": "SDS DIRECT CONTRACT WITH TIE LTD", "page": 1},
        ],
        "reviews": {"6.9.2": {"by": "analyst", "at": "2026-10-18T10:00:00Z"}},
        "partial_reasons": ["coverage_lost_to_budget"],
        "research_audit": {"candidates": [{}] * 184},
    },
}

TCPR_ORDER = ["tcW", "gridSpan", "vMerge", "tcBorders", "shd", "noWrap", "tcMar", "vAlign"]
PPR_ORDER = ["pStyle", "keepNext", "pBdr", "shd", "tabs", "spacing", "ind", "jc", "rPr", "sectPr"]


def _render():
    return build_event_chronology_docx(
        job=JOB, project_name="Edinburgh Tram Inquiry (public)",
        generated=datetime(2026, 10, 18, tzinfo=timezone.utc),
    )


def test_events_follow_date_order_with_review_status():
    overview, events = build_events(JOB["result"])
    assert overview == ["MUDFA diversions ran late across the network."]
    assert [(e.event_id, e.ref, e.date_label) for e in events] == [
        ("EV-0001", "6.9.2", "19 Sep 2005"),
        ("EV-0002", "6.9.3", "25 Nov 2007"),
        ("EV-0003", "6.9.4", "Date not established"),
    ]
    assert events[0].reviewed and not events[1].reviewed
    assert events[1].delay == ("N", "Notice")
    assert (events[1].from_party, events[1].to_party) == ("tie limited", "Infraco")
    assert format_date("1992-06", "month", long=True) == "June 1992"


def test_document_follows_the_template_layout():
    doc = Document(io.BytesIO(_render()))
    text = "\n".join(p.text for p in doc.paragraphs)
    for heading in ("Event Chronology", "1.  Basis and limitations", "2.  Records reviewed",
                    "3.  Conventions", "4.  Chronology of events", "5.  Delay event schedule",
                    "Appendix A  —  Source register", "End of chronology."):
        assert heading in text
    assert "Of 3 entries, 1 has been verified" in text
    assert "evidence budget" in text

    orientations = [s.orientation for s in doc.sections]
    assert orientations == [WD_ORIENT.PORTRAIT, WD_ORIENT.PORTRAIT, WD_ORIENT.LANDSCAPE]

    cover = {row.cells[0].text: row.cells[1].text for row in doc.tables[0].rows}
    assert cover["Events recorded"] == "3 (EV-0001 to EV-0003)"
    assert cover["Document reference"] == "CHR-009"
    assert "184 documents considered" in cover["Knowledge base"]

    chronology = next(t for t in doc.tables if t.rows[0].cells[0].text == "Event ID" and len(t.columns) == 7)
    assert [c.text for c in chronology.rows[0].cells] == [
        "Event ID", "Date", "Event", "Record type", "Source reference", "From / To", "Rec."]
    assert chronology.rows[1].cells[6].text == "C"
    assert chronology.rows[2].cells[6].text == "C\nU"
    assert "CEC00793598" in chronology.rows[2].cells[4].text
    assert "▲ Notice" in chronology.rows[2].cells[2].text

    register = doc.tables[-1]
    assert [r.cells[0].text for r in register.rows[1:]] == ["CEC00793598", "PBH00035962"]
    assert register.rows[1].cells[5].text == "EV-0002"


def test_xml_is_in_the_order_word_requires():
    with zipfile.ZipFile(io.BytesIO(_render())) as package:
        names = [n for n in package.namelist() if n.startswith("word/") and n.endswith(".xml")]
        for name in names:
            root = ElementTree.fromstring(package.read(name))
            for tag, order in (("tcPr", TCPR_ORDER), ("pPr", PPR_ORDER)):
                for props in root.iter(W + tag):
                    seen = [c.tag.replace(W, "") for c in props if c.tag.replace(W, "") in order]
                    assert seen == sorted(seen, key=order.index), (name, tag, seen)
